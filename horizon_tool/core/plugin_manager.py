"""Discover, read, hash, and templatize Master Prompt plugin files."""
from __future__ import annotations

import hashlib
import re
import shutil
from datetime import datetime
from pathlib import Path

import docx  # python-docx

# Readable plugin formats. Encrypted ``.plugin`` files are intentionally NOT
# discovered yet: their in-memory decryption arrives in Phase 8, and listing
# them now would let the GUI open a file that read_plugin_text cannot read.
SUPPORTED_SUFFIXES = {".txt", ".md", ".docx"}

# Plugins editable in-tool (PL-07). ``.docx``/``.plugin`` are not: a ``.docx``
# selection is redirected to "Mở file" (Word); ``.plugin`` encryption (SEC-08)
# arrives in Phase 8.
HISTORY_DIRNAME = "_history"
EDITABLE_SUFFIXES = {".txt", ".md"}


def list_plugins(folder: Path) -> list[Path]:
    """List readable plugin files directly in a folder (non-recursive).

    Only files with a supported suffix are returned; directories (including the
    ``_history`` backup dir) are excluded because they are not files. Returns an
    empty list if the folder does not exist.
    """
    if not folder.exists():
        return []
    return [
        p for p in folder.iterdir()
        if p.is_file() and p.suffix.lower() in SUPPORTED_SUFFIXES
    ]


def _bold_wrap(text: str) -> str:
    """Wrap the non-space content of a run in Markdown bold.

    Bold delimiters must hug the text (``**word**``); surrounding spaces would
    make CommonMark treat the asterisks as literal, so leading/trailing spaces
    are kept outside the markers. Whitespace-only runs are returned unchanged.
    """
    stripped = text.strip()
    if not stripped:
        return text
    prefix = text[: len(text) - len(text.lstrip())]
    suffix = text[len(text.rstrip()):]
    return f"{prefix}**{stripped}**{suffix}"


def docx_to_markdown(path: Path) -> str:
    """Convert a .docx to Markdown-ish text.

    Heading styles become '#'-prefixed lines (level clamped to 1-6); bold runs
    become **bold** so ChatGPT can see the prompt's structure. One paragraph
    per line.
    """
    document = docx.Document(str(path))
    lines: list[str] = []
    for para in document.paragraphs:
        style = (para.style.name or "").lower() if para.style else ""
        if style.startswith("heading"):
            match = re.search(r"\d+", style)
            level = min(int(match.group()) if match else 1, 6)
            lines.append("#" * level + " " + para.text)
            continue
        parts: list[str] = []
        for run in para.runs:
            text = run.text
            if run.bold:
                text = _bold_wrap(text)
            parts.append(text)
        lines.append("".join(parts) if parts else para.text)
    return "\n".join(lines)


def read_plugin_text(path: Path) -> str:
    """Read a plugin file to text. (.plugin decryption handled in Phase 8.)"""
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8")
    if suffix == ".docx":
        return docx_to_markdown(path)
    raise ValueError(f"Unsupported or not-yet-supported plugin type: {path.suffix}")


def plugin_hash(text: str) -> str:
    """SHA-256 hex digest of plugin content, for versioning in logs/reports."""
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def backup_plugin(path: Path) -> Path | None:
    """Copy the current plugin file into plugins/_history/ with a timestamp.

    Returns the backup path, or None if the source doesn't exist yet (nothing
    to back up for a brand-new file).
    """
    path = Path(path)
    if not path.exists():
        return None
    history = path.parent / HISTORY_DIRNAME
    history.mkdir(parents=True, exist_ok=True)
    # Microseconds so two saves in the same second don't overwrite each other.
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    backup = history / f"{path.stem}.{stamp}{path.suffix}"
    shutil.copy2(path, backup)
    return backup


def save_plugin_text(path: Path, text: str) -> None:
    """Save edited plugin text, backing up the previous version first (PL-07).

    Only plaintext plugins (.txt/.md) are editable in-tool; .docx/.plugin raise.
    """
    path = Path(path)
    if path.suffix.lower() not in EDITABLE_SUFFIXES:
        raise ValueError(f"Không sửa được trong tool (chỉ .txt/.md): {path.suffix}")
    backup_plugin(path)
    path.write_text(text, encoding="utf-8")


def apply_variables(text: str, variables: dict[str, str]) -> str:
    """Replace {VAR} placeholders with values; unknown placeholders untouched.

    Values are assumed not to contain ``{OTHER_KEY}`` placeholder syntax (they
    are simple settings like "15s"), so the sequential replacements do not
    re-substitute each other.
    """
    for key, value in variables.items():
        text = text.replace("{" + key + "}", value)
    return text
