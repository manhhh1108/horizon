import docx
import pytest

from horizon_tool.core.plugin_manager import (
    list_plugins, read_plugin_text, plugin_hash, apply_variables,
)


def test_list_plugins_finds_supported(tmp_path):
    (tmp_path / "v11.txt").write_text("prompt", encoding="utf-8")
    (tmp_path / "v12.md").write_text("prompt", encoding="utf-8")
    # list_plugins never opens the file, so a stub .docx is fine for discovery.
    (tmp_path / "v13.docx").write_text("x", encoding="utf-8")
    (tmp_path / "notes.pdf").write_text("ignore", encoding="utf-8")
    (tmp_path / "_history").mkdir()  # excluded

    names = sorted(p.name for p in list_plugins(tmp_path))
    assert names == ["v11.txt", "v12.md", "v13.docx"]


def test_list_plugins_missing_folder_returns_empty(tmp_path):
    assert list_plugins(tmp_path / "nope") == []


def test_read_txt_plugin(tmp_path):
    p = tmp_path / "p.txt"
    p.write_text("Line one\nLine two", encoding="utf-8")
    assert read_plugin_text(p) == "Line one\nLine two"


def test_read_docx_plugin_to_markdown(tmp_path):
    p = tmp_path / "p.docx"
    d = docx.Document()
    d.add_heading("MAIN TITLE", level=1)
    para = d.add_paragraph()
    run = para.add_run("bold part")
    run.bold = True
    d.add_paragraph("plain text")
    d.save(p)

    text = read_plugin_text(p)
    assert "# MAIN TITLE" in text
    assert "**bold part**" in text
    assert "plain text" in text


def test_docx_bold_run_with_surrounding_spaces(tmp_path):
    # A bold run "bold word " must not become "**bold word **" (broken in
    # CommonMark); the trailing space stays outside the bold markers.
    p = tmp_path / "spaces.docx"
    d = docx.Document()
    para = d.add_paragraph()
    para.add_run("before ")
    run = para.add_run("bold word ")
    run.bold = True
    para.add_run("after")
    d.save(p)

    text = read_plugin_text(p)
    assert "before **bold word** after" in text


def test_docx_high_heading_level_is_clamped(tmp_path):
    p = tmp_path / "deep.docx"
    d = docx.Document()
    d.add_heading("DEEP", level=9)  # Word caps at 9; markdown must clamp to 6
    d.save(p)

    text = read_plugin_text(p)
    assert "###### DEEP" in text  # exactly 6 '#', not 9


def test_read_plugin_text_rejects_unsupported(tmp_path):
    p = tmp_path / "secret.plugin"
    p.write_bytes(b"encrypted-bytes")
    with pytest.raises(ValueError):
        read_plugin_text(p)


def test_plugin_hash_stable_and_sensitive():
    h1 = plugin_hash("abc")
    h2 = plugin_hash("abc")
    h3 = plugin_hash("abd")
    assert h1 == h2 and h1 != h3
    assert len(h1) == 64  # sha256 hex


def test_apply_variables_replaces_known_and_leaves_rest():
    text = "Length {VIDEO_DURATION}, quality {VIDEO_QUALITY}, keep {OTHER}"
    out = apply_variables(text, {"VIDEO_DURATION": "15s", "VIDEO_QUALITY": "1080p"})
    assert out == "Length 15s, quality 1080p, keep {OTHER}"


def test_apply_variables_no_placeholder_is_verbatim():
    text = "No variables here."
    assert apply_variables(text, {"VIDEO_DURATION": "15s"}) == text
