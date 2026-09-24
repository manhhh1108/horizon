# horizon_tool/tests/test_word_builder.py
import docx

from horizon_tool.core.section_parser import ParsedSections
from horizon_tool.core.word_builder import build_document


def _parsed():
    return ParsedSections(sections={
        "full_story": "CHAPTER ONE — The Door\n\nHe stood. **The truth was a lie.**",
        "key_scenes": "Scene 1. Scene 2.",
        "image_9x16": "A vertical shot.",
        "thumbnail_16x9": "A key art.",
        "video_prompt": "SHOT 1 — 3s\nSHOT 2 — 3s",
        "fb_title": "his sleeve tore open. what happened next stopped everyone.",
        "fb_description": "A gripping story.",
        "story_teaser": "He never expected it.",
        "hashtags": "#story #viral #drama #twist #fyp",
    }, missing=[])


def _read(path):
    return docx.Document(str(path))


def test_builds_headings_and_bold(tmp_path):
    out = tmp_path / "1.docx"
    build_document(_parsed(), out)
    doc = _read(out)
    styles = [p.style.name for p in doc.paragraphs]
    texts = [p.text for p in doc.paragraphs]
    assert "Heading 1" in styles
    assert any(t == "FULL STORY" for t in texts)           # section title
    assert any(s == "Heading 2" and "CHAPTER ONE" in t
               for s, t in zip(styles, texts))             # chapter heading
    # a real bold run exists
    assert any(r.bold for p in doc.paragraphs for r in p.runs)


def test_facebook_title_uppercase_and_bold(tmp_path):
    out = tmp_path / "2.docx"
    build_document(_parsed(), out)
    doc = _read(out)
    title_paras = [p for p in doc.paragraphs
                   if p.text.startswith("HIS SLEEVE TORE OPEN")]
    assert title_paras
    p = title_paras[0]
    assert p.text == p.text.upper()                        # all uppercase
    assert all(r.bold for r in p.runs if r.text.strip())   # bold


def test_video_prompt_is_monospace(tmp_path):
    out = tmp_path / "3.docx"
    build_document(_parsed(), out)
    doc = _read(out)
    runs = [r for p in doc.paragraphs for r in p.runs if "SHOT 1" in r.text]
    assert runs and runs[0].font.name == "Consolas"


def test_missing_sections_still_builds(tmp_path):
    out = tmp_path / "4.docx"
    parsed = ParsedSections(sections={"full_story": "Just a story."}, missing=[
        "key_scenes", "image_9x16", "thumbnail_16x9", "video_prompt",
        "fb_title", "fb_description", "story_teaser", "hashtags"])
    build_document(parsed, out)
    doc = _read(out)
    assert any(p.text == "FULL STORY" for p in doc.paragraphs)
