import docx

from horizon_tool.core.input_reader import extract_ordinal, read_docx_file, read_script_content


def test_extract_ordinal_plain_number():
    assert extract_ordinal("1.txt") == 1


def test_extract_ordinal_with_words():
    assert extract_ordinal("Kịch bản 12.docx") == 12


def test_extract_ordinal_uses_last_number():
    assert extract_ordinal("2024_scene_5.txt") == 5


def test_extract_ordinal_none_when_absent():
    assert extract_ordinal("intro.txt") is None


def test_read_docx_file(tmp_path):
    p = tmp_path / "5.docx"
    d = docx.Document()
    d.add_paragraph("Dòng một")
    d.add_paragraph("Dòng hai")
    d.save(p)
    assert read_docx_file(p) == "Dòng một\nDòng hai"


def test_read_script_content_dispatches_by_suffix(tmp_path):
    t = tmp_path / "3.txt"
    t.write_text("nội dung txt", encoding="utf-8")
    assert read_script_content(t) == "nội dung txt"
