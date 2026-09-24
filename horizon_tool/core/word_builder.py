"""Build the output Word document from parsed sections."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Pt

from horizon_tool.core.section_parser import SECTION_ORDER, ParsedSections

# Display heading for each section (canonical Master-Prompt names).
SECTION_TITLES = {
    "full_story": "FULL STORY",
    "key_scenes": "KEY SCENES + CONTINUITY NOTE",
    "image_9x16": "IMAGE PROMPT — 9:16",
    "thumbnail_16x9": "THUMBNAIL PROMPT — 16:9",
    "video_prompt": "VIDEO AI PROMPT",
    "fb_title": "FACEBOOK TITLE",
    "fb_description": "FACEBOOK VIDEO DESCRIPTION",
    "story_teaser": "STORY TEASER",
    "hashtags": "HASHTAGS",
}

_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def _add_markdown_runs(paragraph, text: str) -> None:
    """Add text to a paragraph, turning **bold** into real bold runs."""
    pos = 0
    for m in _BOLD_RE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        run = paragraph.add_run(m.group(1))
        run.bold = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _render_story(doc, body: str) -> None:
    for line in body.split("\n"):
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.upper().startswith("CHAPTER "):
            doc.add_heading(stripped, level=2)
        else:
            _add_markdown_runs(doc.add_paragraph(), line)


def _render_monospace(doc, body: str) -> None:
    for line in body.split("\n"):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Pt(18)
        run = paragraph.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(10)


def _render_generic(doc, body: str) -> None:
    for line in body.split("\n"):
        if line.strip():
            _add_markdown_runs(doc.add_paragraph(), line)


def build_document(parsed: ParsedSections, path: Path) -> None:
    """Write parsed sections to a .docx with the required formatting.

    Calibri 11 body; section titles as Heading 1; chapter titles as Heading 2;
    Facebook title uppercase + bold on its own paragraph; Video AI Prompt in a
    monospace (Consolas) indented block. No cover, TOC, header, footer, or page
    numbers (python-docx adds none by default). Missing sections are skipped.
    """
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    for key in SECTION_ORDER:
        if key not in parsed.sections:
            continue
        doc.add_heading(SECTION_TITLES[key], level=1)
        body = parsed.sections[key]
        if key == "full_story":
            _render_story(doc, body)
        elif key == "fb_title":
            paragraph = doc.add_paragraph()
            run = paragraph.add_run(body.strip().upper())
            run.bold = True
        elif key == "video_prompt":
            _render_monospace(doc, body)
        else:
            _render_generic(doc, body)

    Path(path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
