# horizon_tool/tests/test_report.py
import openpyxl

from horizon_tool.core.report import ReportWriter, STATUS_REJECTED


def test_report_writes_rows_and_header(tmp_path):
    path = tmp_path / "report.xlsx"
    rw = ReportWriter(path)
    rw.add_row(ordinal=1, input_filename="1.txt", plugin_name="v11",
               plugin_hash="abc123", account="chatgpt_1", word="Xong",
               img_9x16="Xong", img_16x9=STATUS_REJECTED, video="Bỏ qua",
               missing_sections="", error="", duration_seconds=42.5)
    rw.save()

    wb = openpyxl.load_workbook(path)
    ws = wb.active
    header = [c.value for c in ws[1]]
    assert "STT" in header and "Video" in header
    row2 = [c.value for c in ws[2]]
    assert 1 in row2
    assert "v11" in row2


def test_report_highlights_rejected(tmp_path):
    path = tmp_path / "report.xlsx"
    rw = ReportWriter(path)
    rw.add_row(ordinal=2, input_filename="2.txt", plugin_name="v11",
               plugin_hash="h", account="chatgpt_1", word="Xong",
               img_9x16=STATUS_REJECTED, img_16x9="Xong", video="Bỏ qua",
               missing_sections="", error="", duration_seconds=1.0)
    rw.save()
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    # at least one cell in the data row has a non-default fill (highlight)
    fills = [ws.cell(row=2, column=c).fill.fgColor.rgb for c in range(1, ws.max_column + 1)]
    assert any(f not in (None, "00000000") for f in fills)
